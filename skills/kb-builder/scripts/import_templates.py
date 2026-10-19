#!/usr/bin/env python3
"""
Import an Excel/CSV file into Solvea template library.

Required columns:
  - name
  - content

Optional columns:
  - status: defaults to ENABLED
  - scene: empty/default for normal templates, voice_sms for SMS templates
  - AUTO, ENGLISH, CHINESE, JAPANESE, FRENCH, GERMAN, SPANISH, ARABIC,
    RUSSIAN, PORTUGUESE, ITALIAN, KOREAN, DUTCH, INDONESIAN, TURKISH,
    ROMANIAN, POLISH, PHILIPPINES

Usage:
  python3 /skills/kb-builder/scripts/import_templates.py --xlsx /mounted/chat-files/user/templates.xlsx --dry-run
  python3 /skills/kb-builder/scripts/import_templates.py --xlsx /mounted/chat-files/user/templates.xlsx --name-col "Topic" --content-col "Message" --dry-run
  python3 /skills/kb-builder/scripts/import_templates.py --xlsx /mounted/chat-files/user/templates.xlsx --skip-duplicates --dry-run
  python3 /skills/kb-builder/scripts/import_templates.py --plan-file /workspace/outputs/kb-builder/template_import_plan_<ts>.json
  python3 /skills/kb-builder/scripts/import_templates.py --xlsx /mounted/chat-files/user/templates.xlsx
"""
from __future__ import annotations

import argparse
import csv
import html
import json
import os
import sys
from pathlib import Path
from typing import Any

sys.path.insert(0, os.path.join(os.environ.get("SOLVEA_LIB_DIR", os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "__lib__"))), "skill-runtime"))
from bootstrap import *  # noqa: F403

from client import SolveaClient

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - runtime image should include openpyxl in sandbox deps
    load_workbook = None


HEADER_ALIASES = {
    "name": {
        "name",
        "title",
        "topic",
        "template",
        "template_name",
        "template name",
        "template title",
        "模板名称",
        "模板名",
        "模版名称",
        "模版名",
        "名称",
        "标题",
        "主题",
    },
    "content": {
        "content",
        "message",
        "body",
        "reply",
        "answer",
        "response",
        "template_content",
        "template content",
        "reply text",
        "模板内容",
        "模板回复",
        "模版内容",
        "模版回复",
        "回复内容",
        "回复话术",
        "话术",
        "答案",
        "内容",
    },
    "status": {"status", "state", "状态", "发布状态"},
    "scene": {"scene", "usage", "scenario", "场景", "适用场景"},
}

LANGUAGE_COLUMNS = {
    "AUTO",
    "ENGLISH",
    "CHINESE",
    "JAPANESE",
    "FRENCH",
    "GERMAN",
    "SPANISH",
    "ARABIC",
    "RUSSIAN",
    "PORTUGUESE",
    "ITALIAN",
    "KOREAN",
    "DUTCH",
    "INDONESIAN",
    "TURKISH",
    "ROMANIAN",
    "POLISH",
    "PHILIPPINES",
}


def normalize_header(value: Any) -> str:
    return str(value or "").strip().lower().replace("-", "_")


def normalize_cell(value: Any) -> str:
    if value is None:
        return ""
    return html.unescape(str(value).strip())


def canonical_header(header: Any) -> str:
    raw = str(header or "").strip()
    normalized = normalize_header(raw)
    for canonical, aliases in HEADER_ALIASES.items():
        if normalized in {normalize_header(alias) for alias in aliases}:
            return canonical
    upper = raw.upper().strip()
    if upper in LANGUAGE_COLUMNS:
        return upper
    return normalized


def load_templates_from_excel(
    path: str | Path,
    sheet: str | None = None,
    *,
    name_col: str | None = None,
    content_col: str | None = None,
    status_col: str | None = None,
    scene_col: str | None = None,
    lang_cols: dict[str, str] | None = None,
    skip_duplicates: bool = False,
    include_stats: bool = False,
) -> list[dict[str, Any]] | tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    path = Path(path)
    rows = _load_rows(path, sheet)
    if not rows:
        if include_stats:
            return [], []
        return []

    raw_headers = [normalize_cell(cell) for cell in rows[0]]
    headers = resolve_headers(
        raw_headers,
        name_col=name_col,
        content_col=content_col,
        status_col=status_col,
        scene_col=scene_col,
        lang_cols=lang_cols,
    )
    templates: list[dict[str, Any]] = []
    seen_names: set[str] = set()
    skipped_duplicates: list[dict[str, Any]] = []
    for row_idx, row in enumerate(rows[1:], start=2):
        values = {
            headers[idx]: normalize_cell(row[idx]) if idx < len(row) else ""
            for idx in range(len(headers))
            if headers[idx]
        }
        name = values.get("name", "")
        content = values.get("content", "")
        if not name and not content:
            continue
        if not name or not content:
            raise ValueError(f"row {row_idx}: name and content are required")
        if name in seen_names:
            if skip_duplicates:
                skipped_duplicates.append({"row": row_idx, "name": name})
                continue
            raise ValueError(f"row {row_idx}: duplicate template name in file: {name}")
        seen_names.add(name)

        multi_content = {"AUTO": values.get("AUTO") or content}
        for language in sorted(LANGUAGE_COLUMNS - {"AUTO"}):
            if values.get(language):
                multi_content[language] = values[language]

        templates.append({
            "name": name,
            "content": content,
            "status": values.get("status") or "ENABLED",
            "scene": values.get("scene") or None,
            "multiContent": multi_content,
        })
    if include_stats:
        return templates, skipped_duplicates
    return templates


def resolve_headers(
    raw_headers: list[str],
    *,
    name_col: str | None = None,
    content_col: str | None = None,
    status_col: str | None = None,
    scene_col: str | None = None,
    lang_cols: dict[str, str] | None = None,
) -> list[str]:
    headers = [canonical_header(cell) for cell in raw_headers]

    overrides = {
        "name": name_col,
        "content": content_col,
        "status": status_col,
        "scene": scene_col,
    }
    for canonical, requested_col in overrides.items():
        if requested_col:
            headers[_require_raw_header(raw_headers, requested_col, f"--{canonical.replace('_', '-')}-col")] = canonical

    for language, requested_col in (lang_cols or {}).items():
        language = language.upper()
        if language not in LANGUAGE_COLUMNS:
            raise ValueError(f"Unsupported language '{language}'. Supported languages: {sorted(LANGUAGE_COLUMNS)}")
        headers[_require_raw_header(raw_headers, requested_col, f"--lang-col {language}")] = language

    if "name" not in headers or "content" not in headers:
        raise ValueError(
            "Cannot detect template name/content columns. "
            "Use --name-col and --content-col after inspecting the uploaded file header. "
            f"Available columns: {raw_headers}"
        )
    return headers


def _require_raw_header(raw_headers: list[str], requested_col: str, arg_name: str) -> int:
    requested = normalize_header(requested_col)
    for idx, header in enumerate(raw_headers):
        if normalize_header(header) == requested:
            return idx
    raise ValueError(f"{arg_name} '{requested_col}' not found. Available columns: {raw_headers}")


def _load_rows(path: Path, sheet: str | None) -> list[list[Any]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as f:
            return list(csv.reader(f))
    if suffix not in {".xlsx", ".xlsm"}:
        raise ValueError("Only .xlsx, .xlsm, and .csv files are supported")
    if load_workbook is None:
        raise RuntimeError("openpyxl is required to read Excel files")
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.active
    return [list(row) for row in ws.iter_rows(values_only=True)]


def import_templates(
    client: SolveaClient,
    *,
    xlsx_path: str | Path | None = None,
    plan_file: str | Path | None = None,
    dry_run: bool,
    sheet: str | None = None,
    name_col: str | None = None,
    content_col: str | None = None,
    status_col: str | None = None,
    scene_col: str | None = None,
    lang_cols: dict[str, str] | None = None,
    skip_duplicates: bool = False,
) -> dict[str, Any]:
    if bool(xlsx_path) == bool(plan_file):
        raise ValueError("Exactly one of xlsx_path or plan_file is required")
    if plan_file:
        templates = load_templates_from_plan(plan_file)
        skipped_duplicates: list[dict[str, Any]] = []
        source_path = str(plan_file)
    else:
        templates, skipped_duplicates = load_templates_from_excel(
            xlsx_path,
            sheet=sheet,
            name_col=name_col,
            content_col=content_col,
            status_col=status_col,
            scene_col=scene_col,
            lang_cols=lang_cols,
            skip_duplicates=skip_duplicates,
            include_stats=True,
        )
        source_path = str(xlsx_path)
    OUTPUTS_DIR.mkdir(parents=True, exist_ok=True)  # noqa: F405
    plan_path: Path | None = None
    if dry_run:
        plan_path = OUTPUTS_DIR / f"template_import_plan_{timestamp()}.json"  # noqa: F405
        plan_path.write_text(json.dumps(templates, ensure_ascii=False, indent=2), encoding="utf-8")

    created = 0
    failed = 0
    created_items: list[dict[str, Any]] = []
    errors: list[dict[str, str]] = []
    if not dry_run:
        for item in templates:
            try:
                response = client.create_template(item)
                template_id = extract_template_id(response)
                created += 1
                created_items.append({"id": template_id, "name": item["name"]})
                suffix = f" id={template_id}" if template_id is not None else ""
                print(f"  [OK] {item['name']}{suffix}")
            except Exception as exc:
                failed += 1
                errors.append({"name": item["name"], "error": str(exc)})
                print(f"  [FAIL] {item['name']}: {exc}", file=sys.stderr)

    result_path: Path | None = None
    if not dry_run:
        result_path = OUTPUTS_DIR / f"template_import_result_{timestamp()}.json"  # noqa: F405

    result = {
        "input": len(templates),
        "created": created,
        "failed": failed,
        "source_path": source_path,
        "plan_path": str(plan_path) if plan_path else None,
        "result_path": str(result_path) if result_path else None,
        "created_items": created_items,
        "skipped_duplicates": skipped_duplicates,
        "errors": errors,
    }
    if result_path:
        result_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return result


def extract_template_id(response: Any) -> Any:
    if not isinstance(response, dict):
        return None
    if response.get("id") is not None:
        return response.get("id")
    data = response.get("data")
    if isinstance(data, dict):
        return data.get("id")
    return None


def load_templates_from_plan(plan_file: str | Path) -> list[dict[str, Any]]:
    path = Path(plan_file)
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, list):
        raise ValueError("template import plan must be a JSON array")
    for idx, item in enumerate(data, start=1):
        if not isinstance(item, dict) or not item.get("name") or not item.get("content"):
            raise ValueError(f"invalid template import plan item at index {idx}")
    return data


def main() -> None:
    parser = argparse.ArgumentParser(description="Import Excel/CSV templates into Solvea template library")
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument("--xlsx", default=None, help="Excel/CSV file path")
    source.add_argument("--plan-file", default=None, help="Import from a dry-run template_import_plan JSON file")
    parser.add_argument("--sheet", default=None, help="Optional Excel sheet name")
    parser.add_argument("--name-col", default=None, help="Column containing template name")
    parser.add_argument("--content-col", default=None, help="Column containing default template content")
    parser.add_argument("--status-col", default=None, help="Optional status column")
    parser.add_argument("--scene-col", default=None, help="Optional scene column")
    parser.add_argument(
        "--lang-col",
        action="append",
        default=[],
        help="Optional language column mapping, e.g. CHINESE=中文话术. Can be repeated.",
    )
    parser.add_argument("--dry-run", action="store_true", help="Parse and write plan without creating templates")
    parser.add_argument(
        "--skip-duplicates",
        action="store_true",
        help="Keep the first row for duplicate template names and skip later duplicates",
    )
    args = parser.parse_args()

    import_templates(
        SolveaClient(),
        xlsx_path=args.xlsx,
        plan_file=args.plan_file,
        sheet=args.sheet,
        dry_run=args.dry_run,
        name_col=args.name_col,
        content_col=args.content_col,
        status_col=args.status_col,
        scene_col=args.scene_col,
        lang_cols=parse_lang_cols(args.lang_col),
        skip_duplicates=args.skip_duplicates,
    )


def parse_lang_cols(values: list[str]) -> dict[str, str]:
    result: dict[str, str] = {}
    for value in values:
        if "=" not in value:
            raise ValueError(f"--lang-col must use LANGUAGE=COLUMN format, got: {value}")
        language, column = value.split("=", 1)
        language = language.strip().upper()
        column = column.strip()
        if not language or not column:
            raise ValueError(f"--lang-col must use LANGUAGE=COLUMN format, got: {value}")
        result[language] = column
    return result


if __name__ == "__main__":
    main()

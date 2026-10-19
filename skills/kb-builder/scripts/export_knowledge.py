#!/usr/bin/env python3
"""
Export Solvea knowledge base to an Excel file with exact directory paths.

The script reads the full directory tree once, fetches knowledge through
getKnowledge, joins each knowledge item by directoryId, and writes one Excel
sheet with dynamic Level columns.

Usage:
  python3 /skills/kb-builder/scripts/export_knowledge.py --all
  python3 /skills/kb-builder/scripts/export_knowledge.py --directory-id 105809
  python3 /skills/kb-builder/scripts/export_knowledge.py --directory-id 105809 --keyword battery
  python3 /skills/kb-builder/scripts/export_knowledge.py --directory-id 105809 --no-children
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

sys.path.insert(0, os.path.join(os.environ.get("SOLVEA_LIB_DIR", os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "__lib__"))), "skill-runtime"))
from bootstrap import *  # noqa: F401,F403

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from client import SolveaClient


UNKNOWN_DIRECTORY = "Unknown Directory"
ILLEGAL_EXCEL_CHARS_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


@dataclass
class DirectoryIndex:
    nodes: dict[int, dict[str, Any]] = field(default_factory=dict)
    paths: dict[int, list[str]] = field(default_factory=dict)
    children: dict[int, list[int]] = field(default_factory=dict)

    def descendant_ids(self, directory_id: int) -> set[int]:
        ids = {int(directory_id)}
        for child_id in self.children.get(int(directory_id), []):
            ids.update(self.descendant_ids(child_id))
        return ids

    def max_depth(self, directory_ids: set[int] | None = None) -> int:
        if directory_ids is None:
            paths = self.paths.values()
        else:
            paths = [self.paths.get(int(directory_id), []) for directory_id in directory_ids]
        return max((len(path) for path in paths), default=1) or 1


def _directory_roots(response: Any) -> list[dict[str, Any]]:
    if isinstance(response, list):
        return response
    if isinstance(response, dict):
        value = response.get("list") or response.get("records") or response.get("data") or []
        return value if isinstance(value, list) else []
    return []


def build_directory_index(response: Any) -> DirectoryIndex:
    index = DirectoryIndex()

    def walk(nodes: list[dict[str, Any]], parent_path: list[str]) -> None:
        for node in nodes:
            directory_id = node.get("id")
            if directory_id is None:
                continue
            directory_id = int(directory_id)
            name = str(node.get("name") or "").strip()
            path = parent_path + ([name] if name else [])
            parent_id = int(node.get("parentId") or 0)

            index.nodes[directory_id] = node
            index.paths[directory_id] = path
            index.children.setdefault(parent_id, []).append(directory_id)
            index.children.setdefault(directory_id, [])

            children = node.get("children") or []
            if isinstance(children, list) and children:
                walk(children, path)

    walk(_directory_roots(response), [])
    return index


def _knowledge_items(response: Any) -> list[dict[str, Any]]:
    if isinstance(response, list):
        return response
    if isinstance(response, dict):
        value = response.get("list") or response.get("records") or response.get("data") or []
        return value if isinstance(value, list) else []
    return []


def _knowledge_total(response: Any, fallback: int) -> int:
    if isinstance(response, dict):
        for key in ("total", "totalCount", "count"):
            value = response.get(key)
            if isinstance(value, int):
                return value
    return fallback


def fetch_knowledge_page(
    client: SolveaClient,
    *,
    page_index: int,
    page_size: int,
    keyword: str,
    knowledge_type: str,
    directory_id: int | None,
) -> Any:
    return client.list_knowledge(
        page_index=page_index,
        page_size=page_size,
        keyword=keyword,
        knowledge_type=knowledge_type,
        directory_id=directory_id,
    )


def fetch_all_knowledge(
    client: SolveaClient,
    *,
    directory_id: int | None,
    keyword: str,
    knowledge_type: str,
    page_size: int,
) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    page = 1
    while True:
        response = fetch_knowledge_page(
            client,
            page_index=page,
            page_size=page_size,
            keyword=keyword,
            knowledge_type=knowledge_type,
            directory_id=directory_id,
        )
        page_items = _knowledge_items(response)
        if not page_items:
            break
        items.extend(page_items)

        total = _knowledge_total(response, len(items))
        if len(items) >= total or len(page_items) < page_size:
            break
        page += 1
    return items


def selected_directory_ids(
    directory_index: DirectoryIndex,
    directory_ids: list[int] | None,
    *,
    include_children: bool,
) -> set[int] | None:
    if not directory_ids:
        return None
    selected: set[int] = set()
    for directory_id in directory_ids:
        directory_id = int(directory_id)
        if include_children:
            selected.update(directory_index.descendant_ids(directory_id))
        else:
            selected.add(directory_id)
    return selected


def normalize_item(item: dict[str, Any]) -> dict[str, Any]:
    return {
        "id": item.get("id", item.get("knowledgeId", item.get("kuId", ""))),
        "directoryId": item.get("directoryId"),
        "question": item.get("question", item.get("title", item.get("name", ""))) or "",
        "answer": item.get("answer", item.get("content", item.get("brief", ""))) or "",
        "type": item.get("type", ""),
        "status": item.get("status", item.get("knowledgeStatus", "")) or "",
    }


def clean_excel_value(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    return ILLEGAL_EXCEL_CHARS_RE.sub("", value)


def collect_knowledge(
    client: SolveaClient,
    directory_index: DirectoryIndex,
    *,
    directory_ids: list[int] | None,
    include_children: bool,
    keyword: str,
    knowledge_type: str,
    page_size: int,
) -> tuple[list[dict[str, Any]], set[int] | None]:
    selected_ids = selected_directory_ids(
        directory_index,
        directory_ids,
        include_children=include_children,
    )
    raw_items: list[dict[str, Any]] = []

    if selected_ids is None:
        raw_items = fetch_all_knowledge(
            client,
            directory_id=None,
            keyword=keyword,
            knowledge_type=knowledge_type,
            page_size=page_size,
        )
    else:
        for directory_id in sorted(selected_ids):
            raw_items.extend(
                fetch_all_knowledge(
                    client,
                    directory_id=directory_id,
                    keyword=keyword,
                    knowledge_type=knowledge_type,
                    page_size=page_size,
                )
            )

    deduped: dict[str, dict[str, Any]] = {}
    for item in raw_items:
        normalized = normalize_item(item)
        key = str(normalized["id"] or f"{normalized['directoryId']}:{normalized['question']}")
        deduped[key] = normalized
    return list(deduped.values()), selected_ids


def _path_for_item(directory_index: DirectoryIndex, item: dict[str, Any]) -> list[str]:
    directory_id = item.get("directoryId")
    if directory_id is None:
        return [UNKNOWN_DIRECTORY]
    try:
        return directory_index.paths.get(int(directory_id), [UNKNOWN_DIRECTORY])
    except (TypeError, ValueError):
        return [UNKNOWN_DIRECTORY]


def write_excel(
    items: list[dict[str, Any]],
    directory_index: DirectoryIndex,
    output_path: Path,
    *,
    selected_ids: set[int] | None,
) -> None:
    paths = [_path_for_item(directory_index, item) for item in items]
    if selected_ids is None:
        max_depth = max((len(path) for path in paths), default=directory_index.max_depth())
    else:
        max_depth = max(
            directory_index.max_depth(selected_ids),
            max((len(path) for path in paths), default=1),
        )
    max_depth = max(max_depth, 1)

    wb = Workbook()
    ws = wb.active
    ws.title = "Knowledge Base"

    headers = [f"Level {i}" for i in range(1, max_depth + 1)]
    headers.extend(["Question", "Answer", "KU ID", "Type", "Status", "Directory ID"])

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    rows = sorted(
        zip(items, paths),
        key=lambda pair: (pair[1], str(pair[0].get("question", "")), str(pair[0].get("id", ""))),
    )
    for row_idx, (item, path) in enumerate(rows, 2):
        padded_path = path + [""] * (max_depth - len(path))
        values = padded_path[:max_depth] + [
            item.get("question", ""),
            item.get("answer", ""),
            item.get("id", ""),
            item.get("type", ""),
            item.get("status", ""),
            item.get("directoryId", ""),
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=clean_excel_value(value))
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=col > max_depth)

    widths = [22] * max_depth + [48, 80, 14, 14, 18, 14]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = ws.cell(row=2, column=max_depth + 1).coordinate
    ws.auto_filter.ref = ws.dimensions

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def write_cache(
    directory_response: Any,
    items: list[dict[str, Any]],
    *,
    directory_ids: list[int] | None,
    include_children: bool,
    keyword: str,
    knowledge_type: str,
) -> None:
    export_tmp = TMP_DIR / "export"
    export_tmp.mkdir(parents=True, exist_ok=True)
    (export_tmp / "directories.json").write_text(
        json.dumps(directory_response, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    (export_tmp / "knowledge_items.json").write_text(
        json.dumps(items, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    (export_tmp / "request.json").write_text(
        json.dumps(
            {
                "directoryIds": directory_ids,
                "includeChildren": include_children,
                "keyword": keyword,
                "type": knowledge_type,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )


def build_directory_stats(items: list[dict[str, Any]], directory_index: DirectoryIndex) -> list[dict[str, Any]]:
    counts: dict[str, dict[str, Any]] = {}
    for item in items:
        directory_id = item.get("directoryId")
        path = _path_for_item(directory_index, item)
        path_text = " / ".join(path)
        entry = counts.setdefault(
            path_text,
            {"path": path_text, "directoryId": directory_id, "count": 0},
        )
        entry["count"] += 1
    return sorted(counts.values(), key=lambda row: (-row["count"], row["path"]))


def export_knowledge(
    client: SolveaClient,
    *,
    output_path: Path | None = None,
    directory_ids: list[int] | None = None,
    include_children: bool = True,
    keyword: str = "",
    knowledge_type: str = "ALL",
    page_size: int = 1000,
    write_tmp_cache: bool = False,
) -> dict[str, Any]:
    directory_response = client.list_directories()
    directory_index = build_directory_index(directory_response)
    items, selected_ids = collect_knowledge(
        client,
        directory_index,
        directory_ids=directory_ids,
        include_children=include_children,
        keyword=keyword,
        knowledge_type=knowledge_type,
        page_size=page_size,
    )

    if output_path is None:
        output_path = OUTPUTS_DIR / f"kb_export_{timestamp()}.xlsx"
    else:
        output_path = Path(output_path)

    write_excel(items, directory_index, output_path, selected_ids=selected_ids)
    directory_stats = build_directory_stats(items, directory_index)
    if write_tmp_cache:
        write_cache(
            directory_response,
            items,
            directory_ids=directory_ids,
            include_children=include_children,
            keyword=keyword,
            knowledge_type=knowledge_type,
        )

    return {
        "output_path": str(output_path),
        "knowledge_count": len(items),
        "directory_count": len(directory_index.nodes),
        "selected_directory_count": len(selected_ids) if selected_ids is not None else None,
        "directory_stats": directory_stats,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Export Solvea knowledge with directory columns")
    parser.add_argument("--all", action="store_true", dest="export_all",
                        help="Export all knowledge. This is the default when no --directory-id is provided.")
    parser.add_argument("--directory-id", action="append", type=int, default=[],
                        help="Directory ID to export. Can be specified multiple times.")
    parser.add_argument("--no-children", action="store_true",
                        help="Only export the specified directory IDs, not their descendants.")
    parser.add_argument("--keyword", default="", help="Keyword filter passed to getKnowledge")
    parser.add_argument("--type", dest="knowledge_type", default="ALL",
                        help="Knowledge type: ALL / QA / FILE / WEBSITE / PRODUCT / LEARN_SELF")
    parser.add_argument("--page-size", type=int, default=1000)
    parser.add_argument("--output", default="", help="Optional output xlsx path")
    args = parser.parse_args()

    output_path = Path(args.output) if args.output else None
    directory_ids = args.directory_id or None
    result = export_knowledge(
        SolveaClient(),
        output_path=output_path,
        directory_ids=directory_ids,
        include_children=not args.no_children,
        keyword=args.keyword,
        knowledge_type=args.knowledge_type,
        page_size=args.page_size,
        write_tmp_cache=True,
    )

    print(f"[export] directories={result['directory_count']}")
    if result["selected_directory_count"] is not None:
        print(f"[export] selected_directories={result['selected_directory_count']}")
    print(f"[export] knowledge={result['knowledge_count']}")
    for row in result["directory_stats"][:20]:
        print(f"[export] directory count={row['count']} id={row['directoryId']} path={row['path']}")
    if len(result["directory_stats"]) > 20:
        print(f"[export] directory ... {len(result['directory_stats']) - 20} more")
    print(f"[export] output={result['output_path']}")


if __name__ == "__main__":
    main()
